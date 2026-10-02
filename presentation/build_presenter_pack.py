"""Presenter pack for Part B (Hasnat) and Part C (Shakil): slide-by-slide scripts, demo cheat sheet,
plain-English glossary and a Q&A bank. Output: Word document."""
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

OUT = sys.argv[1] if len(sys.argv) > 1 else "ENT-HYBRID-NET_Part_B_C_Presenter_Pack.docx"
NAVY, TEAL, GREY, AMBER = RGBColor(0x0B, 0x1F, 0x3A), RGBColor(0x0F, 0x83, 0x8F), RGBColor(0x55, 0x60, 0x70), RGBColor(0xB4, 0x6A, 0x00)

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.0)
    s.top_margin = s.bottom_margin = Cm(1.8)
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
for name, size, color in (("Heading 1", 18, NAVY), ("Heading 2", 14, TEAL), ("Heading 3", 12, NAVY)):
    h = doc.styles[name]
    h.font.name, h.font.size, h.font.color.rgb, h.font.bold = "Calibri", Pt(size), color, True


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_fill)
    tcPr.append(sh)


def para(text, bold=False, italic=False, color=None, size=None, after=4):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(after)
    return p


def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        if isinstance(it, tuple):
            p.add_run(it[0]).bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)
        p.paragraph_format.space_after = Pt(2)


def say_box(text):
    """The words to SAY, in a shaded box."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]
    shade(c, "EAF6F7")
    c.paragraphs[0].add_run("SAY:  ").bold = True
    c.paragraphs[0].runs[0].font.color.rgb = TEAL
    first = True
    for chunk in text.strip().split("\n\n"):
        p = c.paragraphs[0] if first else c.add_paragraph()
        p.add_run(chunk.strip())
        p.paragraph_format.space_after = Pt(4)
        first = False
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def code(lines):
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    shade(c, "1E293B")
    for i, l in enumerate(lines):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        r = p.add_run(l)
        r.font.name, r.font.size = "Consolas", Pt(9.5)
        r.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def slide(num, title, on_slide, script, point=None, demo=None, time=None):
    doc.add_heading(f"Slide {num}: {title}" + (f"   ({time})" if time else ""), level=3)
    para("On the slide: " + on_slide, italic=True, color=GREY, size=10)
    say_box(script)
    if point:
        para("Point at: " + point, color=AMBER, size=10)
    if demo:
        para("Live demo (only if asked or if time allows):", bold=True, size=10)
        code(demo)


def qa(q, a):
    p = doc.add_paragraph()
    p.add_run("Q: ").bold = True
    r = p.add_run(q)
    r.bold = True
    r.font.color.rgb = NAVY
    p.paragraph_format.space_after = Pt(1)
    p2 = doc.add_paragraph()
    p2.add_run("A: ").bold = True
    p2.add_run(a)
    p2.paragraph_format.left_indent = Cm(0.6)
    p2.paragraph_format.space_after = Pt(7)


# ============================================================ cover
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("ENT-HYBRID-NET  |  MN521 Network Automation")
r.font.size, r.font.color.rgb = Pt(12), TEAL
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Presenter Pack: Part B and Part C")
r.bold, r.font.size, r.font.color.rgb = True, Pt(26), NAVY
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Slide-by-slide scripts, demo cheat sheet, plain-English glossary and likely questions with answers")
r.italic, r.font.size, r.font.color.rgb = True, Pt(12), GREY

tbl = doc.add_table(rows=4, cols=3)
tbl.style = "Table Grid"
rows = [("Presenter", "Part", "Slides"), ("Shah Abul Hasnat Chowdhury (MIT261576)", "Part B: Ansible automation", "8 to 13  (about 6 min)"),
        ("Shakil Ahammed Shahan (MIT261626)", "Part C: Terraform, AWS and Git", "14 to 17  (about 5 min)"),
        ("Shared close (suggested)", "Testing, challenges, conclusion", "18 Shakil, 19 Hasnat, 20 Shakil")]
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        c = tbl.rows[i].cells[j]
        c.text = v
        if i == 0:
            shade(c, "0B1F3A")
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            c.paragraphs[0].runs[0].bold = True
doc.add_paragraph()
doc.add_heading("How to use this pack (read first, 2 minutes)", level=2)
bullets([
    ("Read the SAY boxes aloud twice. ", "They are written to be spoken. Don't learn them word for word: learn the three numbers and the one idea on each slide."),
    ("Learn the cheat card on the next page. ", "If you forget everything else, those numbers and sentences will carry you."),
    ("Questions: ", "answer in one or two sentences, then stop. If you don't know, say: \"Good question. The short answer is X; the details are in section 4 of our report.\" Never guess a number."),
    ("Demos are optional. ", "Only run one if the marker asks. The commands are in the cheat sheet, and every result is already a screenshot in 07_Evidence_Screenshots."),
    ("Hand-overs: ", "Hasnat starts on slide 8 after Jahidul. Shakil starts on slide 14. Use the hand-over line at the end of each part."),
])

# ============================================================ cheat card
doc.add_page_break()
doc.add_heading("Cheat card: numbers and one-liners to memorise", level=1)
card = doc.add_table(rows=1, cols=2)
card.style = "Table Grid"
left, right = card.rows[0].cells
shade(left, "EAF6F7")
shade(right, "FFF4E0")
left.paragraphs[0].add_run("PART B  (Hasnat)").bold = True
for l in ["10 devices automated: 4 routers + 6 switches", "8 roles, 10 playbooks, all run by site.yml", "14 / 14 compliance rules pass on every device",
          "changed=0 on the second run = idempotent", "0 failed, 0% packet loss in reachability tests", "Vault = AES-256, vault password never in Git",
          "Control node: ANSIBLE-SRV 10.30.100.10 (NetOps VLAN 100)", "Connection: network_cli over SSH v2 (cisco.ios collection)",
          "Day-2 change: CHG-0001, VLAN 45 MARKETING at Branch", "Health thresholds: CPU < 80% (5 min), memory >= 10% free",
          "Tags: --tags ntp,syslog | --tags health | --list-tags"]:
    left.add_paragraph(l, style="List Bullet")
right.paragraphs[0].add_run("PART C  (Shakil)").bold = True
for l in ["Region ap-southeast-2 (Sydney), 2 Availability Zones", "VPC 172.20.0.0/16 (no overlap with on-prem 10.0.0.0/8)",
          "Public 172.20.1.0/24 + 172.20.2.0/24  |  Private 172.20.11.0/24 + 172.20.12.0/24", "25 resources created in about 30 seconds",
          "Bastion: SSH only from admin /32. App server: no public IP", "Workflow: init > fmt > validate > plan > apply > verify > destroy",
          "After destroy: 0 billable resources in all 17 regions", "No NAT gateway on purpose: it is not free tier",
          "t3.micro, Amazon Linux 2023, IMDSv2, encrypted gp3 disks", "Git: staged commits + CHG-0001; state, keys, tfvars git-ignored"]:
    right.add_paragraph(l, style="List Bullet")
doc.add_paragraph()
para("The one idea of the whole project (both of you can say this):", bold=True, color=NAVY)
say_box("The desired state of the network and the cloud lives in code. That code sits in Git, it is applied automatically, and it is tested every time it runs. That is what Infrastructure as Code means.")

# ============================================================ glossary
doc.add_heading("Plain-English glossary (what the words mean)", level=1)
doc.add_heading("Part B words", level=2)
gl_b = [
    ("Ansible: ", "a tool that logs in to devices over SSH and configures them from a written description. Nothing is installed on the routers ('agentless')."),
    ("Control node: ", "the computer that runs Ansible. Ours is ANSIBLE-SRV, a Linux container in the Data Centre (10.30.100.10)."),
    ("Inventory (hosts.yml): ", "the list of devices and their groups: by role (routers, switches) and by site (HQ, Branch, DC, Core)."),
    ("group_vars / host_vars: ", "YAML files holding the data: VLANs, IP addresses, ACLs. This is our 'single source of truth'."),
    ("Playbook: ", "a file that says which devices get which roles. site.yml runs playbooks 00 to 08 in order."),
    ("Role: ", "a reusable package of tasks for one job, for example 'routing' or 'baseline'. We have 8."),
    ("Task / module: ", "one step, for example 'push these config lines'. cisco.ios.ios_config is the module that pushes Cisco config."),
    ("Jinja2 template: ", "a config file with blanks. Ansible fills the blanks from the YAML data, e.g. the OSPF template."),
    ("Idempotent: ", "running it again changes nothing if the device is already correct. Our proof: second run shows changed=0."),
    ("Handler: ", "a task that only runs when something changed. Ours saves the config to NVRAM (write memory)."),
    ("Check mode (--check --diff): ", "a dry run. It shows what WOULD change without touching any device."),
    ("Tags: ", "labels on tasks so you can run just one part, e.g. --tags ntp runs only the NTP tasks."),
    ("Ansible Vault: ", "encrypts secrets (passwords, SNMP keys) with AES-256 so they can live safely in Git."),
    ("network_cli: ", "how Ansible talks to network devices: an SSH session to the CLI, like a human would."),
    ("Compliance audit: ", "the 08 playbook checks 14 security rules against every running config and writes a report."),
]
bullets(gl_b)
doc.add_heading("Part C words", level=2)
gl_c = [
    ("Terraform: ", "a tool that builds cloud infrastructure from code (.tf files). You describe what you want; it works out how."),
    ("Declarative: ", "you write the end result ('one VPC, four subnets'), not the steps."),
    ("Provider: ", "a plug-in that talks to a cloud's API. We use the AWS provider (plus tls and local for the key file)."),
    ("State file: ", "Terraform's memory of what it built. It compares state, code and the real cloud to plan changes. Never in Git (it can hold secrets)."),
    ("plan / apply / destroy: ", "plan = preview, apply = build, destroy = delete everything Terraform built."),
    ("VPC: ", "your own private network inside AWS. Ours is 172.20.0.0/16."),
    ("Availability Zone (AZ): ", "a separate data centre inside a region. Two AZs = more resilience."),
    ("Public / private subnet: ", "public subnets have a route to the Internet Gateway; private ones do not."),
    ("Internet Gateway (IGW): ", "the door between the VPC and the Internet."),
    ("Route table: ", "the rules that say where traffic goes. Public: 0.0.0.0/0 to IGW. Private: local only."),
    ("Security group: ", "a firewall around each server. Ours allow SSH to the bastion from one IP, and to the app server only from the bastion."),
    ("Bastion host (jump box): ", "the only server you can SSH into from outside. You hop through it to reach private servers."),
    ("ProxyJump (ssh -J): ", "SSH through the bastion to the private server in one command."),
    ("NAT gateway: ", "lets private servers reach the Internet outbound. Costs money, so it is off by default in our code."),
    ("IMDSv2: ", "a safer version of the EC2 metadata service that needs a session token. It blocks a common credential-theft attack (SSRF)."),
    ("Key pair: ", "SSH keys for logging in. Terraform generated an ED25519 key; the private part stayed on our PC and was deleted after destroy."),
]
bullets(gl_c)

# ============================================================ PART B script
doc.add_page_break()
doc.add_heading("Part B script: Shah Abul Hasnat Chowdhury (slides 8 to 13)", level=1)
para("Target time: about 6 minutes. Speak slowly. One idea per slide.", italic=True, color=GREY)

slide(8, "Part B title", "Part B: Enterprise network automation with Ansible. 'The network as code: YAML data, Jinja2 templates and eight roles run from ANSIBLE-SRV (10.30.100.10).'",
      """Thanks, Jahidul. Jahidul showed you the network we designed. My part is how we stopped configuring it by hand.

We describe the whole network as data in YAML files, we turn that data into Cisco configuration with templates, and Ansible pushes it to all ten devices from one Linux server in the Data Centre, ANSIBLE-SRV.""", time="20 s")

slide(9, "How the automation works", "Pipeline: inventory + vars, Jinja2 templates, cisco.ios modules, 10 devices. Cards: Vault, Handlers, Tags, Idempotent.",
      """Here is the flow from left to right.

First, the inventory and variables. hosts.yml lists our ten devices and groups them two ways: by role, routers and switches, and by site, HQ, Branch and Data Centre. The group_vars and host_vars files hold the real data: VLANs, IP addresses, ACLs. That is our single source of truth.

Second, Jinja2 templates turn that data into Cisco configuration, for example the OSPF and interface config.

Third, the cisco.ios modules log in over SSH version 2, compare what we want with what is running, and push only the difference.

The four boxes underneath are the quality features. Vault encrypts our passwords with AES-256. A handler saves the config to NVRAM only when something actually changed, and never during a dry run. Tags let us run just one piece, for example only NTP and Syslog. And it is idempotent: run it twice, and the second run changes nothing.""",
      point="the arrow from 'Inventory + vars' to '10 devices', then the four cards", time="1 min 15 s",
      demo=["cd /root/ansible && source venv/bin/activate", "ansible-inventory --graph", "ansible-playbook playbooks/02_baseline.yml --ask-vault-pass --tags ntp,syslog"])

slide(10, "Playbooks, roles and tags", "Table: playbooks 00 to 08, their role, what they automate and their tags.",
      """This table is the whole pipeline. site.yml runs these in order.

We start with 00, which discovers every device and writes an inventory. 01 backs up every running config before we touch anything, which is our safety net.

Then we build and secure: 02 baseline sets the hostname, SSH version 2, the management ACL, NTP, Syslog, DNS, SNMPv3 and the login banner. 03 does the VLANs and trunks, 04 the routing and OSPF, 05 DHCP, DNS and NAT, and 06 the security ACLs.

Finally we check ourselves: 07 verifies the network works, and 08 audits security. Every playbook has tags, so we can run, say, only the verify checks, or only the banner.""",
      point="read the left column top to bottom, then the Tags column", time="1 min",
      demo=["ansible-playbook playbooks/site.yml --list-tags"])

slide(11, "Monitoring: eight automated checks", "Eight numbered checks: backup, interfaces, routing, VLANs, ACLs, CPU and memory, connectivity, inventory.",
      """The brief asked for at least six monitoring tasks. We automated eight.

We back up every config. We check that loopbacks and point-to-point links are up. We check every OSPF neighbour is FULL and the default route is there. We check every switch has its VLANs and trunks. We read the ACL hit counters. We check CPU and memory: the five-minute CPU average must be under 80 percent, and at least 10 percent of memory must be free. We ping the Data Centre servers and the Internet. And we record each device's model and IOS version in an inventory file.

The important point: every check is an assertion. If anything is wrong, the playbook fails loudly instead of passing quietly.""",
      point="go 1 to 8 quickly; slow down on 6 (CPU and memory) and the last sentence", time="1 min",
      demo=["ansible-playbook playbooks/07_verify.yml --ask-vault-pass --tags health", "ansible-playbook playbooks/07_verify.yml --ask-vault-pass"])

slide(12, "Results: compliance and idempotency", "Two screenshots (compliance 14/14, second run changed=0), stats 14/14 and 0% packet loss, NTP note.",
      """Here are the results, straight from Ansible's own output.

On the left, the compliance audit: all 14 security rules pass on all ten devices. Things like SSH version 2 only, no Telnet, the management ACL on the VTY lines, the banner, NTP, Syslog, SNMPv3, and no HTTP server.

On the right, idempotency: we ran the whole configuration pipeline a second time and every device reports changed equals zero. That proves our code describes the network exactly, so it is safe to run again and again.

Reachability tests had zero percent packet loss. To be honest about one thing: NTP is partial. The routers reach the NTP master, but the emulator's clocks drift too much for IOS to declare them synchronised. On real hardware this would work.""",
      point="'14/14' then 'changed=0' in the screenshots", time="1 min")

slide(13, "Day-2 change: VLAN 45 Marketing", "Five steps (edit YAML, commit, dry run, apply, verify) and a screenshot of R2-BR after the change.",
      """Finally, a real change, the kind a network team makes every week. The Branch needed a new Marketing VLAN.

Step one, we added two lines of YAML for VLAN 45. Step two, we committed that to Git as change CHG-0001, so it can be reviewed and rolled back. Step three, a dry run with check and diff showed exactly what would change, without touching any device. Step four, the pipeline applied only that difference. Step five, the devices confirm it: the screenshot shows R2 with the new sub-interface, the DHCP pool BR-MARKETING and OSPF area 20.

Nobody typed a single command on a router. That is the value of automation.

Now Shakil will show how we did the same thing for the cloud with Terraform.""",
      point="the five numbered circles, then the screenshot", time="1 min",
      demo=["ansible-playbook playbooks/03_switching.yml playbooks/04_routing.yml playbooks/05_services.yml -l site_branch --check --diff --ask-vault-pass"])

# ============================================================ PART C script
doc.add_page_break()
doc.add_heading("Part C script: Shakil Ahammed Shahan (slides 14 to 17)", level=1)
para("Target time: about 5 minutes.", italic=True, color=GREY)

slide(14, "Part C title", "Part C: Cloud Infrastructure as Code with Terraform. 'A secure two-AZ AWS VPC with a bastion host and a private EC2 instance, deployed, verified and destroyed from code.'",
      """Thanks, Hasnat. Hasnat automated the on-premises network. I did the same for the cloud.

Using Terraform, we built a secure network in AWS from code, tested it, and then removed it completely, so it costs nothing.""", time="15 s")

slide(15, "AWS architecture", "Architecture diagram and five points: VPC, subnets, bastion, app server, hardening.",
      """This is what Terraform builds, in the Sydney region.

One VPC, 172.20.0.0/16. We chose that range so it never overlaps with our on-premises 10.0.0.0 network, which keeps a future VPN possible.

It spans two Availability Zones. Each zone has a public subnet and a private subnet. The public subnets route to the Internet Gateway; the private ones only route inside the VPC.

Security is the main design idea. Only one server, the bastion, can be reached from the Internet, and only on SSH port 22 from one admin IP address. The application server has no public IP at all, and it accepts SSH only from the bastion's security group. We also enforce IMDSv2 and encrypted disks.

We left the NAT gateway off on purpose. It isn't free tier, so the private server has no Internet access, by design.""",
      point="the bastion in the public subnet, the red SSH arrow, then the app server in the private subnet", time="1 min 30 s")

slide(16, "Terraform workflow and deployment", "Seven workflow steps, screenshot of terraform apply, screenshot of SSH ProxyJump, stats 25 and 0.",
      """This is the Terraform workflow we followed.

init downloads the AWS provider. fmt and validate check the code. plan shows what will be built: 25 resources. We saved that plan, and apply built exactly it, in about 30 seconds. The screenshot on the left shows the outputs, including the bastion's public IP and a ready-made SSH command.

Then we verified it for real. On the right, I SSH into the bastion, it reaches the Internet, and then I jump through it to the private server with ProxyJump. The private server can't reach the Internet, which is exactly what we designed.

Finally, terraform destroy removed all 25 resources. We then checked every AWS region separately with the AWS CLI: zero billable resources left.""",
      point="the workflow pills left to right, then '25', then '0'", time="1 min 30 s",
      demo=["cd terraform", "terraform init", "terraform validate", "terraform plan        # do NOT run apply live unless the marker asks"])

slide(17, "Version control with Git", "git log screenshot, four bullets (staged commits, CHG-0001, secrets, backups) and an Infrastructure as Code definition box.",
      """Everything lives in one GitHub repository.

We committed the Terraform code in stages, providers, then network, then security, then compute, then outputs, so you can see the history. Hasnat's VLAN change is its own commit, CHG-0001, so it can be reviewed or reverted.

Secrets never go into Git: Terraform state, private keys and variable files are ignored, and passwords are encrypted with Vault.

This is what Infrastructure as Code means: the infrastructure is described in versioned files, every change is a commit, and the results are tested automatically.""",
      point="the git log screenshot, then the IaC box", time="1 min")

doc.add_heading("Shared closing slides (suggested split)", level=2)
slide(18, "Testing summary  [Shakil]", "Two tables of tests with Pass, NTP Partial.",
      """We tested in three ways: show commands on the devices, real tests from the end-host PCs, and the Ansible verify and compliance checks after every change. Everything passed except NTP, which is partial because of the emulator's clock drift, as Hasnat explained. Every row has a screenshot in our evidence folder.""", time="30 s")
slide(19, "Challenges and lessons learned  [Hasnat]", "Four challenge cards with problem and fix.",
      """We hit real problems. The simulated PCs couldn't get DHCP from Cisco, so we used small Linux containers. The GNS3 VM froze under Hyper-V, so we wrote a recovery script. Cisco hides default settings, so we tuned our roles until a second run changed nothing. And downloads failed through the lab, so we installed Ansible offline. The lesson: building tests into the automation gave us confidence after every run.""", time="45 s")
slide(20, "Conclusion  [Shakil]", "What we delivered, next steps, Thank you.",
      """To conclude: we built a secure, hierarchical network, we manage it as code with Ansible, and we built and tested a secure AWS environment with Terraform, all tracked in Git. Next we would add a CI/CD pipeline, NetBox as the source of truth, and a site-to-site VPN to AWS. Thank you. We're happy to take questions.""", time="30 s")

# ============================================================ Demo cheat sheet
doc.add_page_break()
doc.add_heading("Live demo cheat sheet", level=1)
para("Only demo if asked. Before the presentation: open GNS3, start all nodes, wait 3 minutes, then run python scripts\\recover_lab.py if anything looks stuck. The vault password is in the team's private notes (not in Git).", size=10)
doc.add_heading("Part B (on ANSIBLE-SRV console)", level=3)
code(["cd /root/ansible && source venv/bin/activate",
      "ansible-inventory --graph                                          # show the inventory groups",
      "ansible-playbook playbooks/site.yml --list-tags                    # every tag in the pipeline",
      "ansible-playbook playbooks/07_verify.yml --ask-vault-pass --tags health      # CPU/memory checks",
      "ansible-playbook playbooks/07_verify.yml --ask-vault-pass          # all 8 checks",
      "ansible-playbook playbooks/08_compliance.yml --ask-vault-pass      # 14-rule audit",
      "ansible-playbook playbooks/02_baseline.yml --ask-vault-pass --tags ntp,syslog --check --diff",
      "ansible-vault view inventory/group_vars/all/vault.yml              # shows it is encrypted"])
doc.add_heading("Part C (Windows terminal, in the terraform folder)", level=3)
code(["terraform version", "terraform init", "terraform fmt -check", "terraform validate", "terraform plan                      # 25 to add, 0 change, 0 destroy",
      "# apply/destroy only if the marker explicitly asks (costs cents, takes ~1 min each)", "git log --oneline --graph            # version history"])
para("If a demo fails: stay calm and say \"The emulator is slow today. Here is the same run captured earlier\", then open the screenshot from 07_Evidence_Screenshots (B10 verify, B12 compliance, B13 idempotency, C05 apply, C10 audit).", italic=True, size=10)

# ============================================================ Q&A
doc.add_page_break()
doc.add_heading("Likely questions and answers", level=1)
para("Answer in one or two sentences. Bold questions are the most likely.", italic=True, color=GREY, size=10)

doc.add_heading("Part B: Ansible (Hasnat)", level=2)
qa_b = [
    ("Why Ansible and not Python scripts or Puppet?", "Ansible is agentless: it only needs SSH, so nothing is installed on the routers. It is declarative YAML that a network team can read, and it has official Cisco IOS modules."),
    ("What does idempotent mean, and how did you prove it?", "Running the same playbook again changes nothing if the device is already correct. We proved it by running the whole pipeline twice: the second run showed changed=0 on all ten devices."),
    ("How does Ansible connect to the routers?", "With the network_cli connection over SSH version 2, using the cisco.ios collection. It logs in as the netadmin user and uses enable for privileged mode."),
    ("Where are the passwords stored?", "In an Ansible Vault file, encrypted with AES-256. The vault password is never committed to Git; we type it with --ask-vault-pass."),
    ("What is a role?", "A reusable package of tasks for one job. We have eight: backup, baseline, vlans, routing, services, acl, verify and compliance."),
    ("What is the difference between group_vars and host_vars?", "group_vars hold data shared by a group, such as all Branch VLANs; host_vars hold data for one device, such as R2's router-ID and its point-to-point link."),
    ("What is a Jinja2 template?", "A configuration file with placeholders. Ansible fills them from the YAML data, for example the OSPF template writes the network statements for each site."),
    ("What does a handler do?", "It runs only when a task reports a change. Ours runs 'write memory' to save the config, and it is skipped in check mode so a dry run never writes NVRAM."),
    ("How do you run only part of the pipeline?", "With tags. For example --tags ntp,syslog runs only those tasks, --tags health runs only the CPU and memory checks, and --list-tags shows them all."),
    ("What does --check --diff do?", "It is a dry run: Ansible shows what would change on each device without changing anything. We used it before the VLAN 45 change."),
    ("What happens if a device is down?", "That host fails and is reported as unreachable or failed in the recap; the others continue. The verify playbook would also fail its assertions, so we notice straight away."),
    ("What do the 14 compliance rules check?", "Things like SSH version 2 only, Telnet disabled, the VTY access class, a local user with secret, enable secret, password encryption, login brute-force protection, the banner, NTP, remote Syslog, archive logging, SNMPv3, no SNMP communities and no HTTP server."),
    ("Why did the first runs show changes every time?", "Cisco hides default commands and rewrites some lines, for example 'logging host' is shown as 'logging'. We changed the roles to push exactly what IOS stores, and then the second run showed changed=0."),
    ("How do you add a new device or site?", "Add it to hosts.yml and give it host_vars, or add a new site group with its VLANs. The same roles then configure it; no new code is needed."),
    ("How are backups done?", "Playbook 01 saves a timestamped running-config for every device and keeps a 'latest' copy in Git, so any change between runs shows up as a normal git diff."),
    ("Why are switches different from routers in your code?", "The switches run older IOS 12.4. Small differences, such as SNMPv3 using DES instead of AES, are handled in group variables, so the same roles work on both."),
    ("Why is NTP only partial?", "The routers reach the NTP master, but the emulated routers' clocks drift so much that IOS refuses to synchronise. The configuration is correct and would work on real hardware."),
    ("What are the CPU and memory thresholds?", "The playbook fails if the five-minute CPU average is 80 percent or more, or if less than 10 percent of processor memory is free."),
    ("How long does the full pipeline take?", "About six to seven minutes for all ten devices in our emulated lab; on real hardware it would be faster."),
]
for q, a in qa_b:
    qa(q, a)

doc.add_heading("Part C: Terraform, AWS and Git (Shakil)", level=2)
qa_c = [
    ("Why Terraform and not CloudFormation?", "Terraform works across many clouds with one language, keeps a state file so it can plan exact changes, and its plan step lets us review before we build. CloudFormation only works for AWS."),
    ("What is the Terraform state file and why isn't it in Git?", "It is Terraform's record of what it built, so it can compare and plan. It can contain sensitive values, so we git-ignore it. For a team we would store it in an encrypted S3 bucket with state locking."),
    ("What is the difference between plan and apply?", "plan only shows what will change; apply makes the change. We saved the plan with -out tfplan and applied that exact file, so what was reviewed is exactly what was built."),
    ("Why two Availability Zones?", "For resilience: if one data centre has a problem, the other still works. The brief also asked for two public and two private subnets."),
    ("Why does the private server have no Internet access?", "The private route table only has the local VPC route. Internet access would need a NAT gateway, which costs money, so we left it off; it can be turned on with one variable, enable_nat_gateway."),
    ("What is a bastion host and why use one?", "A single hardened server that is the only way in from the Internet. Admins SSH to it, then jump to private servers. It reduces the attack surface to one port on one server."),
    ("How is the bastion secured?", "Its security group allows SSH only from our admin IP address as a /32. It uses key-based login, IMDSv2 and an encrypted disk."),
    ("Why does the app server's rule reference the bastion's security group instead of an IP?", "Because if the bastion is rebuilt and gets a new IP, the rule still works. It trusts the role, not the address."),
    ("What is a security group vs a route table?", "A route table decides where traffic goes; a security group is a stateful firewall that decides what traffic is allowed in or out of an instance."),
    ("How did you create the SSH key?", "Terraform generated an ED25519 key with the tls provider, uploaded the public key to AWS and saved the private key locally with restricted permissions. It is git-ignored and was deleted on destroy."),
    ("What is IMDSv2?", "Version 2 of the EC2 instance metadata service. It requires a session token, which blocks attacks that trick a server into leaking its credentials (SSRF)."),
    ("How do you know nothing is still costing money?", "After terraform destroy, we used the AWS CLI to check all 17 regions for instances, volumes, Elastic IPs, NAT gateways, VPCs, key pairs and snapshots. All were zero."),
    ("How much did it cost?", "Almost nothing: two t3.micro instances for about 20 minutes, which is within the free tier or a few cents at most."),
    ("What are input variable validations?", "Rules in variables.tf. For example admin_cidr must be a valid CIDR and cannot be 0.0.0.0/0, so nobody can accidentally open SSH to the world."),
    ("Why did you commit the Terraform code in stages?", "To show proper version control: providers, then network, then security, then compute, then outputs. Each commit is a small, reviewable step."),
    ("What is the .terraform.lock.hcl file?", "It pins the exact provider versions, so every team member gets the same AWS provider and the same result."),
    ("How would this connect to the on-premises network?", "With a site-to-site IPsec VPN between R4 or the Data Centre and an AWS VPN gateway. We chose non-overlapping ranges (172.20.0.0/16 vs 10.0.0.0/8) so this is possible."),
    ("What does terraform destroy do?", "It deletes everything Terraform created, in the right dependency order, and empties the state file."),
]
for q, a in qa_c:
    qa(q, a)

doc.add_heading("General and tricky questions (either presenter)", level=2)
qa_g = [
    ("What is Infrastructure as Code in one sentence?", "Managing infrastructure through versioned, machine-readable files instead of manual commands, so changes are repeatable, reviewable and testable."),
    ("What would you improve next?", "A CI/CD pipeline that runs lint and --check on every Git commit, NetBox as the source of truth, remote Terraform state with locking, and moving the OSPF keys into Vault."),
    ("What was the hardest part?", "Making the automation idempotent on real Cisco IOS, and keeping the emulated lab stable. Both are in our challenges slide."),
    ("Is your Git repository public?", "It is private; we have given the marker access (make sure this is true before you say it)."),
    ("Who did what?", "Jahidul led the network design, Hasnat the Ansible automation and Shakil the Terraform and AWS work. We all tested and wrote the report together."),
    ("Can you show it running now?", "Yes. Offer the safe read-only demos from the cheat sheet: --list-tags, 07_verify with --tags health, or terraform plan. If the lab is slow, show the captured screenshots."),
]
for q, a in qa_g:
    qa(q, a)

doc.save(OUT)
print("wrote", OUT)
